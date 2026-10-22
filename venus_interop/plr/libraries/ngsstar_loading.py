"""Hamilton's NGSStar framework, NGSStarLoading (`Library\\NGSStarLoading\\NGSStarLoading.hsi`): what to load.

A kit method declares, step by step, how much of each reagent one sample takes
(`Define_VolumePerWell`; a master mix's components with `Define_VolumePerWell_MastermixComponent`),
then asks for each reagent's containers (`Calculate_TotalVolumeAndSequenceToAspirate`): how many
tubes or troughs of the given type, how much in each (samples x volume x (1 + dead %) + the
container's dead volume, split so a container holds at most its volume), the positions they stand
on (a sequence of distinct containers, and one position per sample to aspirate from), and, for a
master mix, how much of each component goes into each of its tubes (the Mastermix_Prep sheet).

The library keeps this in an Excel workbook (`RunData.xls`, sheets RunData and Mastermix_Prep)
through Venus's file object with SQL; here the sheets are `Sheet`s with the same columns, text
matched without case as Jet's Excel driver does. Container sizes, dead volumes and minimum counts
are the method's arguments, as in Venus (Avenio passes them as literals).

Kept as the library has them, though they look unintended:
- with a minimum number of containers that the first sizing misses, sizing repeats with a smaller
  maximum volume, and every pass appends its containers' volumes to the volume list returned (and
  the list the Mastermix_Prep rows read), so the list is longer than the containers;
- the display texts (DB_ReagentNameANDLabwareType) keep one decimal fewer than they round to;
- the site list (DB_SiteID, and DB_PositionID for troughs) is built onto a variable the call never
  assigns before, so it starts with that variable's value: "0", if an unassigned variable is
  integer 0 (CONVENTIONS §11); the library empties it only after a master mix.
A component's own containers (the library's other branch) are not ported: Avenio never asks.
"""
from __future__ import annotations

import math
import pathlib
from typing import Any, Optional

from venus_interop.methods import sequences as sq
from venus_interop.plr.operator import RunAborted

LIBRARY = "NGSSTARLOADING"  # the Venus namespace ported

REAGENT, MASTERMIX, COMPONENT = "Reagent", "Mastermix", "MastermixComponent"  # ARW_Global::ReagentType
RUNDATA = ["DB_VolumePerWell", "DB_DeadVolumeLabware", "DB_DeadVolumePercent", "DB_TotalVolume", "DB_CarrierName",
           "DB_SiteID", "DB_Units", "DB_LabwareType", "DB_LabwareName", "DB_PositionID", "DB_ReagentName",
           "DB_MasterMixComponentName", "DB_ReagentType", "DB_PrepareMastermixOnDeck", "DB_Ratio",
           "DB_ReagentNameANDLabwareType"]

class LabwareType:
  """ARW_Global::LabwareType: the names the loading texts show for a container type."""
  HamiltonTrough50mL = "50mL Trough"
  Tube5mL = "tube 4ml Sarstedt"
  Sarstedt_Tube_2mL = "vial 2mL Sarstedt"
  ScrewCap_Tube_1500uL = "Vial 1.5ml Roche (original tube)"
  Sarstedt_Tube_1500uL = "Vial 1.5ml Eppendorf"
  Sarstedt_Tube_500uL = "Vial 0.5ml Sarsdtedt"
  ScrewCap_Tube_500uL = "Vial 0.5ml Roche (original tube)"
  MIDI_Plate = "MIDI Plate"
  HSP_Plate = "HSP Plate"
  Comfortlid = "Comfortlid"
  TIP50 = "50uL Tips"
  TIP300 = "300uL Tips"
  TIP1000 = "1000uL Tips"
  TIPAdapter = "Tip Adapter"


PLT_TIP1000 = "PLT_CAR_2TIP_2PCR_MTP"  # ARW_Global::CarrierType: the carrier with 1000 uL tips and plates
PLT_STACK = "PLT_CAR_L4ST_0001"  # the plate stacker
TIP50_01, TIP50_02 = "TIP_CAR_480BC_TIP_50F_PINK_0001", "TIP_CAR_480BC_TIP_50F_PINK_0002"
TIP300_01, TIP300_02 = "TIP_CAR_480BC_TIP_STF_YELLOW_0001", "TIP_CAR_480BC_TIP_STF_YELLOW_0002"
TIP_SLOTS = {LabwareType.TIP50: 10, LabwareType.TIP300: 10, LabwareType.TIP1000: 2}  # ARW_Global::TipLocations
STACK_UNITS = {LabwareType.MIDI_Plate: 4, LabwareType.HSP_Plate: 4, LabwareType.Comfortlid: 8}  # per stack
TIPS = (50, 300, 1000)  # the tip types, in the library's order (its 1, 2, 3)
TIP_NAMES = {50: LabwareType.TIP50, 300: LabwareType.TIP300, 1000: LabwareType.TIP1000}
# A tip type's carriers to unload and load: both of a pair on an NGS STAR
TIP_CARRIERS = {50: (TIP50_01, TIP50_02), 300: (TIP300_01, TIP300_02), 1000: (PLT_TIP1000,)}
NO_RELOAD = " No reload of this tip type necessary"
PREVENTIVE = " Preventive reload of this tip type necessary"
COVER_OPEN = "Front cover open - close door to proceed run."

MASTERMIX_PREP = ["Mastermix_ID", "Component_ID", "Ratio", "MastermixTube_ID", "MastermixTube_LabwareID",
                  "VolumeMastermixPerTube", "VolumeComponentPerTube", "SourceComponentTube_ID",
                  "SourceComponentTube_LabwareID", "SourceVolumeComponentPerTube"]


def _same(a, b) -> bool:
  return a.casefold() == b.casefold() if isinstance(a, str) and isinstance(b, str) else a == b


class Sheet:
  """A workbook sheet as the library's SQL uses it: rows of column -> value (None: an empty cell)."""

  def __init__(self, columns: list[str]):
    self.columns, self.rows = columns, []

  def where(self, column: str, value) -> list[dict]:
    return [r for r in self.rows if _same(r[column], value)]

  def update(self, column: str, value, **cells) -> None:
    for r in self.where(column, value):
      r.update(cells)

  def append(self, **cells) -> None:
    self.rows.append({c: cells.get(c) for c in self.columns})


def mth_round(x: float, places: int) -> float:
  """HSLMthLibImpl's MthRound: half up, from the floor."""
  whole = math.floor(x)
  r = whole + math.floor(10.0 ** places * (x - whole) + 0.5) / 10.0 ** places
  return math.floor(r) if places == 0 else r


def _text(v) -> str:
  """A number as HSL's string concatenation writes it (as the interpreter models it; unverified)."""
  if isinstance(v, float):
    return str(int(v)) if v.is_integer() else repr(v)
  return str(v)


def _volume_text(v: float) -> str:
  """_GenerateString / _Reformat_ValuesForDialog: mL above 3000 uL; rounded to 3 places (2 above
  3000 uL or 3 mL), then cut at the point's index plus the places, which keeps one fewer."""
  value, units, top = (v / 1000.0, "mL", 3) if v > 3000 else (v, "uL", 3000)
  places = 3 if value <= top else 2
  s = _text(mth_round(value, places))
  i = s.find(".")
  return (s[:i + places] if i != -1 else s) + units



def _dialog_number(v: float) -> str:
  """_Reformat_ValuesForDialog on a float in uL: rounded to 3 places (2 above 3000), the text cut
  at the point's index plus the places; no unit, no mL (unlike _volume_text)."""
  places = 3 if v <= 3000 else 2
  s = _text(mth_round(v, places))
  i = s.find(".")
  return s[:i + places] if i != -1 else s

class NGSStarLoading:
  def __init__(self, deck, ngsstar: bool = True, star=None, operator=None, report=None, planner=None,
               faithful: bool = False):
    """`deck`: `seq(name)` (live sequences) and `layout` (the layout file, for carriers and sites).
    `ngsstar`: an NGS STAR (8 channels) rather than a Clinical NGS STARlet (4). For loading the
    deck: `star` (its autoload and front cover), `operator` (an operator.Operator), `report` (the
    run report, ngsstar_report.RunReport, one in memory when none) and `planner` (STAR Tip Tools,
    for the tip counters; one on `deck` when none)."""
    from venus_interop.plr.libraries.ngsstar_report import RunReport
    from venus_interop.plr.libraries.star_tip_tools import StarTipTools

    self.v, self.star, self.operator, self.faithful = deck, star, operator, faithful
    self.sample_ids: list[str] = []  # the input worklist's, for the sample tracking
    self.report = report if report is not None else RunReport()
    self.planner = planner or StarTipTools(deck, faithful=faithful)
    self.DATA_HandOverDataToSMT_NGSStarLoading(ngsstar=ngsstar)

  def DATA_HandOverDataToSMT_NGSStarLoading(self, report_file: str = "", platform: int = 1, kit_provider: str = "",
                                            kit_name: str = "", mph_used: int = 0, lock_unlock: int = 0,
                                            counter50: str = "", counter300: str = "", counter1000: str = "",
                                            tips50=None, tips300=None, tips1000=None, check_barcodes: int = 0,
                                            ngsstar: Optional[bool] = None) -> None:
    """The run's start: the platform, a fresh RunData workbook, the trough bookkeeping emptied, the
    tip counters' names and copies of the tip sequences (the library keeps its own). Whether to
    lock the front cover around loading, and (`check_barcodes`) to read the 1000 uL carrier's
    other labware's barcodes too. The report file is the RunReport given."""
    self.ngsstar = platform == 1 if ngsstar is None else ngsstar
    self.run_data, self.mastermix_prep = Sheet(RUNDATA), Sheet(MASTERMIX_PREP)
    self.rgt_total_volume: list[float] = []  # t_arr_flt_RGT_TotalVolume, ..._RGT_ReagentName
    self.rgt_reagent_name: list[str] = []
    self.rgt_sequences: list[sq.Sequence] = []  # t_arr_seq_RGT_SequencesToAspirate
    self.last_mastermix_name = ""
    self.lock_unlock, self.check_barcodes = bool(int(lock_unlock)), bool(int(check_barcodes))
    self.counters = {50: counter50, 300: counter300, 1000: counter1000}
    self.tips = {t: self.v.seq(s).copy() for t, s in zip(TIPS, (tips50, tips300, tips1000)) if s is not None}
    platform_name = "NGSStar" if self.ngsstar else "ClinicalNGSStarlet"
    folder = ("NGSStarMPH" if int(mph_used) else "NGSStarChannel") if self.ngsstar else "ClinicalNGSStarlet"
    self.pictures = (f"C:\\Program Files (x86)\\HAMILTON\\{platform_name}\\{kit_provider}\\{kit_name}"
                     f"\\Library\\NGSStarLoading\\{folder}Pictures\\")  # t_str_LoadingPicturesFileDirectory
    self.carriers_to_load: list[str] = []  # t_arr_strCarrierName_Distinct
    self.totals: dict[str, int] = {}

  # -- what one sample takes -----------------------------------------------------------------

  def Define_VolumePerWell(self, liquid: str, volume_per_sample: float) -> float:
    """Add a use of `liquid` to its row (a new Reagent row when it has none)."""
    rows = self.run_data.where("DB_ReagentName", liquid)
    if rows:
      total = sum((r["DB_VolumePerWell"] or 0) for r in rows) + volume_per_sample
      self.run_data.update("DB_ReagentName", liquid, DB_VolumePerWell=total)
    else:
      self.run_data.append(DB_VolumePerWell=volume_per_sample, DB_ReagentName=liquid, DB_ReagentType=REAGENT)
    return volume_per_sample

  def Define_VolumePerWell_MastermixComponent(self, mastermix: str, component: str, volume_per_sample: float) -> float:
    """Make `mastermix` a master mix (ratio 1) and add a component row, "<component> - <mastermix>",
    with its share of the master mix's volume per sample."""
    self.run_data.update("DB_ReagentName", mastermix, DB_Ratio=1.0, DB_ReagentType=MASTERMIX)
    mm_volume = self.run_data.where("DB_ReagentName", mastermix)[0]["DB_VolumePerWell"]
    self.run_data.append(DB_VolumePerWell=volume_per_sample, DB_ReagentName=f"{component} - {mastermix}",
                         DB_ReagentType=COMPONENT, DB_MasterMixComponentName=component,
                         DB_Ratio=volume_per_sample / mm_volume)
    return volume_per_sample

  # -- what to load --------------------------------------------------------------------------

  def _carrier_and_site(self, labware: str) -> tuple[str, str]:
    """_GetCarrierIDandSiteID_FromLabID: the labware's carrier (itself when it has none) and site."""
    p = next((q for q in self.v.layout.labware if q.id == labware), None)
    carrier = p.fields.get("Template", "") if p is not None else ""
    return (labware if carrier == "default" else carrier), (p.fields.get("SiteId", "") if p is not None else "")

  def _channel_positions(self) -> list[str]:
    """A trough's positions, one per channel: 1-8, or 2, 4, 6, 8 for four channels."""
    return [str(i) for i in (range(1, 9) if self.ngsstar else range(2, 9, 2))]

  def Calculate_TotalVolumeAndSequenceToAspirate(self, ML_STAR, reagent: str, samples: int, labware_type: str,
                                                 max_volume: float, dead_volume_labware: float,
                                                 dead_volume_percent: float, template: str, starting_point: int,
                                                 min_containers: int):
    """The containers for `reagent` on `template`'s positions (from `starting_point`, or its current
    one when 0; the template is moved past them). Returns (the distinct containers, a position per
    sample, the volume per container); RunData gets the totals, a master mix's tubes Mastermix_Prep."""
    return self._calculate(reagent, samples, labware_type, max_volume, dead_volume_labware, dead_volume_percent,
                           template, starting_point, min_containers, fixed=False)

  def Calculate_TotalVolumeAndSequenceToAspirate_FixedTubeNumber(self, ML_STAR, reagent: str, samples: int,
                                                                 labware_type: str, max_volume: float,
                                                                 dead_volume_labware: float, dead_volume_percent: float,
                                                                 template: str, starting_point: int, containers: int):
    """The same with exactly `containers` containers, whatever their volume. Kept as the library has
    it: a count that doesn't divide the samples adds the remainder to them, once, and each container
    is sized for the whole part (10 samples in 4 tubes: 12, 3 each; 11: 14, still 3 each, 12)."""
    return self._calculate(reagent, samples, labware_type, max_volume, dead_volume_labware, dead_volume_percent,
                           template, starting_point, containers, fixed=True)

  def _calculate(self, reagent, samples, labware_type, max_volume, dead_volume_labware, dead_volume_percent, template,
                 starting_point, min_containers, fixed):
    tpl = self.v.seq(template)
    if samples <= min_containers:
      min_containers = samples
    found = self.run_data.where("DB_ReagentName", reagent)
    if not found:
      if any(_same(r["DB_MasterMixComponentName"], reagent) and r["DB_ReagentType"] == COMPONENT
             for r in self.run_data.rows):
        raise NotImplementedError(f"NGSStarLoading: containers for the master mix component {reagent!r}")
      return sq.Sequence("distinct", []), sq.Sequence("samples", []), []  # "Reagent ... DOES NOT EXIST"
    row = found[-1]
    if not row["DB_VolumePerWell"]:
      return sq.Sequence("distinct", []), sq.Sequence("samples", []), []
    names, types, volumes = [reagent], [row["DB_ReagentType"]], [row["DB_VolumePerWell"]]
    per_sample_mix = row["DB_VolumePerWell"]
    components: list[str] = []
    is_reagent = row["DB_ReagentType"] != MASTERMIX
    if not is_reagent:  # the master mix's components: rows naming it, of the component type
      for r in self.run_data.rows:
        if reagent in (r["DB_ReagentName"] or "") and r["DB_ReagentType"] == COMPONENT:
          components.append(r["DB_MasterMixComponentName"])
          names.append(r["DB_ReagentName"])
          volumes.append(r["DB_VolumePerWell"])
          types.append(r["DB_ReagentType"])
    per_container_all, ratios, tube_ids, tube_labware = [], [], [], []
    out_volumes: list[float] = []
    distinct = samples_seq = None
    containers = total_for_mix = 0
    dead_lab, dead_pct, type_id = dead_volume_labware, dead_volume_percent, labware_type
    carrier, site_ids = "", "0"  # the never-assigned site list: integer 0 (see the module docstring)
    define = True
    for k, (name, kind, volume) in enumerate(zip(names, types, volumes), start=1):
      if kind == COMPONENT:
        define = False
      factor = dead_pct / 100.0 + 1.0
      if kind == COMPONENT:  # Calculate_TotalVolume_MastermixComponent
        per_sample = volume
        needed = volume * total_for_mix / per_sample_mix
        ratio = volume / per_sample_mix
        ratios.append(ratio)
        self.run_data.update("DB_ReagentName", name, DB_TotalVolume=needed,
                             DB_ReagentNameANDLabwareType=f"{name}-->[{_volume_text(needed)}]",
                             DB_Ratio=mth_round(ratio, 2))
        update = False
      else:
        self.last_mastermix_name = reagent
        if fixed and min_containers:
          samples += samples % min_containers
        per_sample, needed, update = volume, samples * volume, True
      needed *= factor
      min_containers = min_containers or 1
      top, passes, required, display = max_volume, 0, 0.0, ""
      while containers < min_containers:
        passes += 1
        room = top - dead_lab
        containers = min_containers if fixed else math.ceil(needed / room)
        most = math.floor(room / (per_sample * factor))
        if most == 0:
          each = needed / containers
          sizes = [each] * containers
          required = needed
        else:
          containers = min_containers if fixed else math.ceil(samples / most)
          base, extra = samples // containers, 0 if fixed else samples % containers
          sizes = [mth_round((base + 1) * per_sample * factor + dead_lab, 1)] * extra
          each = mth_round(base * per_sample * factor + dead_lab, 1)
          sizes += [each] * (containers - extra)
          required = 0.0
          for size in sizes:
            required += size
        for size in sizes:
          if k == 1:
            out_volumes.append(size)
          per_container_all.append(size)
          if "RGT" in (tpl.current_item() or ("", ""))[0]:
            self.rgt_total_volume.append(size)
            self.rgt_reagent_name.append(name)
        display = f"{reagent}-->{containers} x {type_id}" + "".join(f" [{_volume_text(s)}]" for s in sizes)
        if containers < min_containers:
          top = each - 0.1
        if passes == 200:  # Venus shows "Cannot define n containers" and goes on
          break
      if define:
        if starting_point:
          tpl.set_current_position(starting_point)
        items, positions, labware_ids = [], "", ""
        for c in range(1, containers + 1):
          lab, pos = tpl.current_item() or ("", "")  # past the end: an empty item, as in Venus
          last = c == containers
          carrier, site = self._carrier_and_site(lab)
          if "RGT" in lab:
            trough = [(lab, p) for p in self._channel_positions()]
            items += trough
            self.rgt_sequences.append(sq.Sequence(f"{lab}_trough", trough))
            pos = trough[-1][1]
            site_ids += site + ("" if last else ",")
            labware_ids += lab + ("" if last else ";")
            positions = site_ids
          else:
            items.append((lab, pos))
            positions += pos + ("" if last else ",")
            site_ids += site + ("" if last else ";")
            labware_ids += lab + ("" if last else ";")
          tube_ids.append(pos)
          tube_labware.append(lab)
          tpl.increment(1)
        distinct = sq.Sequence("distinct", list(items))
        trough_last = bool(items) and "RGT" in items[-1][0]
        samples_seq = sq.Sequence("samples", list(items) if trough_last else
                                  [items[i % len(items)] for i in range(samples)] if items else [])
      if update:
        self.run_data.update("DB_ReagentName", name, DB_DeadVolumeLabware=dead_lab, DB_DeadVolumePercent=dead_pct,
                             DB_TotalVolume=required, DB_Units=containers, DB_LabwareType=type_id,
                             DB_ReagentNameANDLabwareType=display, DB_CarrierName=carrier, DB_SiteID=site_ids,
                             DB_PositionID=positions, DB_LabwareName=labware_ids)
      if k == 1:  # the master mix done: its components follow without dead volumes
        total_for_mix = required
        dead_lab = dead_pct = 0
        containers, type_id, carrier, site_ids = 0, "", "", ""
    distinct = distinct or sq.Sequence("distinct", [])
    if not is_reagent:  # a row per master-mix tube and component
      for i in range(distinct.total):
        for comp, ratio in zip(components, ratios):
          self.mastermix_prep.append(
            Mastermix_ID=reagent, Component_ID=comp, Ratio=ratio, MastermixTube_ID=tube_ids[i],
            MastermixTube_LabwareID=tube_labware[i], VolumeMastermixPerTube=per_container_all[i],
            VolumeComponentPerTube=per_container_all[i] * ratio, SourceComponentTube_ID="",
            SourceComponentTube_LabwareID="", SourceVolumeComponentPerTube="")
    return distinct, samples_seq or sq.Sequence("samples", []), out_volumes

  def Calculate_ReformatSequenceToMatchNumberOfSamples(self, ML_STAR, starting_point: int, template: str,
                                                       samples: int) -> sq.Sequence:
    """A position per sample, cycling as many template positions as its current labware has
    (counted before moving to `starting_point`, or its current one when 0, so a late start runs
    past the end: an empty item, then position 1 again); a trough (an "RGT" labware) gives its
    channel positions and is kept as a trough to aspirate from. The template itself is not moved."""
    tpl = self.v.seq(template).copy()
    count = tpl.position_count_for_current_labware()
    if starting_point:
      tpl.set_current_position(starting_point)
    items = []
    for _ in range(count):
      lab, pos = tpl.current_item() or ("", "")
      if "RGT" in lab:
        trough = [(lab, p) for p in self._channel_positions()]
        items += trough
        self.rgt_sequences.append(sq.Sequence(f"{lab}_trough", trough))
      else:
        items.append((lab, pos))
      tpl.increment(1)
    return sq.Sequence("samples", [items[i % len(items)] for i in range(samples)] if items else [])

  # -- labware a process uses ----------------------------------------------------------------

  def Add_LabwareForCurrentProcess(self, ML_STAR, labware_type: str, labware_name: str, units: int, template: str,
                                   position_in_sequence=0) -> None:
    """RunData rows for `units` of a labware on `template` (plates, lids, tips), for the loading
    display: a row per carrier the template's labware stands on, or, when the row exists, its unit
    count raised (on every row of that type and name). Kept as the library has them: a carrier's
    site is its own distinct sites' entry at the carrier's index; whether the labware is on a tube
    carrier (a site "...T-...": the position is the row's) is judged by the last labware; a tip
    adapter's row keeps the previous carrier's position."""
    if not units:
      return
    tpl = self.v.seq(template)
    found = [r for r in self.run_data.rows if _same(r["DB_LabwareType"], labware_type)
             and _same(r["DB_ReagentName"], labware_name) and _same(r["DB_LabwareName"], template)]
    if found:
      cells = {"DB_Units": (found[-1]["DB_Units"] or 0) + units}
      if "T-" in self._carrier_and_site((tpl.current_item() or ("", ""))[0])[1]:  # not checked against the HSL
        tokens = f"{found[-1]['DB_PositionID'] or ''},{position_in_sequence}".split(",")
        cells["DB_PositionID"] = ",".join(sorted(tokens))
      for r in self.run_data.rows:
        if _same(r["DB_LabwareType"], labware_type) and _same(r["DB_ReagentName"], labware_name):
          r.update(cells)
      return
    labware = list(dict.fromkeys(lab for lab, _ in tpl.items))
    places = [self._carrier_and_site(lab) for lab in labware]
    tube = bool(places) and "T-" in places[-1][1]
    position = ""
    for k, carrier in enumerate(dict.fromkeys(c for c, _ in places)):
      sites = list(dict.fromkeys(s for c, s in places if c == carrier))
      site = sites[k] if k < len(sites) else ""  # ElementAt past the end: CONVENTIONS Â§11
      if tube:
        position = str(position_in_sequence)
      elif "TIP" in carrier and PLT_TIP1000 not in carrier:
        slots = TIP_SLOTS.get(labware_type, 0)  # an unassigned variable otherwise
        if labware_type == LabwareType.TIPAdapter:
          site = "Tip Adapter"
        else:
          position = f"1 - {slots}"
      else:
        position = ",".join(sites)
      self.run_data.append(DB_VolumePerWell=0, DB_DeadVolumeLabware=0, DB_DeadVolumePercent=0, DB_TotalVolume=0,
                           DB_CarrierName=carrier, DB_SiteID=site, DB_Units=units, DB_LabwareType=labware_type,
                           DB_LabwareName=template, DB_PositionID=position, DB_ReagentName=labware_name,
                           DB_MasterMixComponentName="", DB_ReagentType="LabwareResource", DB_PrepareMastermixOnDeck=0,
                           DB_ReagentNameANDLabwareType=labware_name if labware_type == labware_name
                           else f"{labware_type} ({labware_name})")

  # -- loading the deck ----------------------------------------------------------------------

  def Evaluate_ResourcesByCarrier(self) -> None:
    """The carriers RunData names (what to load), and the totals the method reads: each tip type's
    units (its last row's), and the plates and lids on the stacker and the ODTC (summed)."""
    rows = self.run_data.rows
    self.carriers_to_load = list(dict.fromkeys(r["DB_CarrierName"] for r in rows if r["DB_CarrierName"]))

    def units(labware_type, carrier=None):
      found = [r for r in rows if _same(r["DB_LabwareType"], labware_type)
               and (carrier is None or _same(r["DB_CarrierName"], carrier))]
      if not found:
        return 0
      return sum(r["DB_Units"] or 0 for r in found) if carrier else (found[-1]["DB_Units"] or 0)

    self.totals = {"50": units(LabwareType.TIP50), "300": units(LabwareType.TIP300),
                   "1000": units(LabwareType.TIP1000), "MIDI": units(LabwareType.MIDI_Plate, PLT_STACK),
                   "HSP": units(LabwareType.HSP_Plate, PLT_STACK),
                   "ComfortLid": units(LabwareType.Comfortlid, "MFX_CAR_ODTC_MTP_0001")}

  def _loading_list(self, carrier: str):
    """What the loading dialog shows for `carrier`, from its RunData rows: (position, quantity,
    description) rows sorted by position; the 1000 uL carrier's positions to read barcodes on; the
    rows' reagent names; the last row. Kept as the library has them: on the 1000 uL carrier a
    tip row's position is every tip position so far ("1,2,"); a stacker's sites each show the
    stack's size or the remainder, not what is left (three sites for 6 plates: 4, 4, 2)."""
    positions, quantities, descriptions, names = [], [], [], []
    plt1000, multiple, sites, per_stack, last = "", 0, 0, 0, None
    for r in self.run_data.where("DB_CarrierName", carrier):
      last = r
      names.append(r["DB_ReagentName"])
      pos = r["DB_PositionID"] or ""
      if carrier == PLT_TIP1000:
        if r["DB_LabwareType"] == LabwareType.TIP1000:
          plt1000 += pos + ","
          positions.append(plt1000)
        else:
          if self.check_barcodes:
            plt1000 += pos + ","
          positions.append(pos)
      else:
        multiple = pos.find(",")
        if not pos:
          positions.append(r["DB_SiteID"] or "")
        elif multiple >= 0 and r["DB_CarrierName"] == PLT_STACK:
          tokens = [t for t in pos.split(",") if t]
          sites = len(tokens)
          positions += tokens
        else:
          sites = 1
          positions.append(pos)
      if carrier == PLT_STACK:
        per_stack = STACK_UNITS.get(r["DB_LabwareType"], per_stack)
      description = r["DB_ReagentNameANDLabwareType"] or ""
      units = r["DB_Units"] or 0
      if "RGT" in carrier:
        quantities.append(_volume_text(r["DB_TotalVolume"] or 0.0))
        descriptions.append(description)
      elif multiple != 0 and carrier == PLT_STACK:
        if units > per_stack * sites:
          raise RunAborted("The number of elements to stack exceeds the total number of positions in "
                           "the stack deck sequence. Review maximum number of elements in stack. Method will abort")
        last_site = units % per_stack or per_stack
        done, per_site = False, 0
        for k in range(1, sites + 1):
          if not done:
            per_site = units if units <= per_stack else (per_stack if k != sites else last_site)
          quantities.append(str(per_site))
          descriptions.append(description)
          done = units - per_site == 0
          if done:
            per_site = 0
      else:
        quantities.append(_text(units))
        descriptions.append(description)
    rows = sorted(f"{_position_text(p)}${q}${d}" for p, q, d in zip(positions, quantities, descriptions))
    table = [tuple(row.split("$")[:3]) for row in rows]
    return table, plt1000[:-1], names, last

  async def _cover_unlocked(self):
    if self.lock_unlock:
      await cover_unlocked(self.star)

  async def _cover_locked(self):
    if self.lock_unlock:
      await cover_locked(self.star, self._ask)

  # -- the input worklist ---------------------------------------------------------------------

  SELECT_BOTTOM = "Press Accept to confirm, Cancel to abort"

  @staticmethod
  def _workflow_text(kind: str, ids, points) -> str:
    """The dialog's top text: "Define Workflow ", then the deactivated points, one a line."""
    text = "Define Workflow "
    for k, p in enumerate(points):
      text += (f"Deactivated {kind} Points:\n\n" if k == 0 else "\n") + f"- [{ids[int(p) - 1]}] "
    return text

  def Display_SelectProcesses_Dialog(self, method_selection, process_ids, combinations, no_start, no_stop):
    """The workflow the operator runs: where to start (the processes not in `no_start`, the first
    of them preselected), then where to stop (the processes the start's combination lists, "1.2.3",
    not in `no_stop`, the first preselected). Returns (start, stop, each process on or off): on
    those of the combination that may be stopped at, then off before the start and after the stop.
    Kept as the library has it: a process of the combination that may not be a stop point keeps
    its earlier state (off, unless it is the start). Cancel aborts the run."""
    ids = [str(i) for i in process_ids]
    n = len(ids)
    no_start, no_stop = [int(i) for i in no_start], [int(i) for i in no_stop]
    starts = [i for i in range(1, n + 1) if i not in no_start]
    start = self._choose(self._workflow_text("Starting", ids, no_start), [ids[i - 1] for i in starts],
                         "Start Process", ids)
    on = [i == start for i in range(1, n + 1)]
    combination = [int(float(x)) for x in str(combinations[start - 1]).split(".") if x]  # StrTokenize, StrIVal
    stops = []
    for i in range(1, n + 1):
      if i in combination:
        if i not in no_stop:
          stops.append(i)
          on[i - 1] = True
      else:
        on[i - 1] = False
    stop = self._choose(self._workflow_text("Stop", ids, no_stop), [ids[i - 1] for i in stops], "Stop Process", ids)
    on = [o and start <= i <= stop for i, o in enumerate(on, start=1)]
    return start, stop, on

  def _choose(self, text: str, options: list[str], group: str, ids: list[str]) -> int:
    """A radio-button dialog: the 1-based index (in `ids`) of the option chosen; the first option
    preselected (an operator answering True takes it); Cancel (None) aborts."""
    chosen = self._ask("choose", text, options=options, title=f"{group} - {self.SELECT_BOTTOM}")
    if chosen is None or chosen is False:
      raise RunAborted(f"{group}: cancelled")
    if chosen is True:
      chosen = options[0]
    return ids.index(chosen) + 1

  def Display_InputWorklist_Dialog(self) -> str:
    """The input worklist chosen by the operator (a file dialog; asked again while none is given;
    the operator's cancel aborts the run, as Venus's abort button does)."""
    trial = 0
    while True:
      trial += 1
      path = self._ask("ask_file", "Please select Input File Worklist", pattern="*")
      if path:
        self.worklist_path = path  # g_str_InputWorklistFilePath
        return path
      self._ask("message", f"Error during previous selection of a worklist Select again a worklist - Trial : {trial}")

  def Retrieve_50mlTrough_ReagentsVolumesAndSequences(self):
    """The reagents in 50 mL troughs, as Calculate sized them: (names, total volumes, sequences)."""
    return list(self.rgt_reagent_name), list(self.rgt_total_volume), list(self.rgt_sequences)

  MASTERMIX_COLUMNS = ["Name", "Volume (uL)", "Labware Type and [Volume per tube]"]
  MASTERMIX_TEXT = "Prepare Mastermix according to instructions below. Press OK to Continue"

  def Display_Mastermix_Dialog(self, log_folder=None):
    """The master mixes to prepare, shown to the operator as a loading list and written as
    `MastermixSummary.csv` (`;` between fields, in `log_folder`; Venus's is HAMILTON\\Logfiles):
    every RunData row that is a master mix or one of its components (a component's name indented
    `---  `, a blank row before each master mix after the first), its total volume (none where the
    labware is "Empty"), and what follows "-->" in its reagent-and-labware text."""
    import csv

    rows, mixes = [], 0
    for r in self.run_data.rows:
      kind = r.get("DB_ReagentType")
      if kind is None or kind in (REAGENT, "LabwareResource"):  # SQL's <> passes no empty cell either
        continue
      if kind == MASTERMIX:
        mixes += 1
        if mixes > 1:
          rows.append(("", "", ""))
      name = f"---  {r.get('DB_MasterMixComponentName', '')}" if kind == COMPONENT else r.get("DB_ReagentName", "")
      combo = str(r.get("DB_ReagentNameANDLabwareType") or "")
      i = combo.find("-->")
      tail = combo[len(combo) - (len(combo) - i - 3):] if len(combo) - i - 3 > 0 else ""  # StrRight
      volume = "" if "Empty" in tail else _dialog_number(float(r.get("DB_TotalVolume") or 0))
      rows.append((name, volume, tail))
    if not rows:
      return
    bottom = "A .csv file will be automatically created at following directory : \nC:\\Program Files (x86)\\HAMILTON\\Logfiles\\"
    self._ask("show_table", self.MASTERMIX_TEXT, title=bottom, columns=self.MASTERMIX_COLUMNS, rows=rows,
              images=[self.pictures + "MasterMix.png"])
    # the dialog's answer is not read: Cancel goes on as OK does (as the library has it)
    if log_folder is not None:
      with (pathlib.Path(log_folder) / "MastermixSummary.csv").open("w", newline="", encoding="utf-8") as f:
        w = csv.writer(f, delimiter=";", quoting=csv.QUOTE_ALL)  # as Venus writes text fields
        w.writerow(["Reagent Name", "Volume Required", "Description"])
        w.writerows(rows)

  WORKLIST_COLUMNS = ("Sample_ID", "Sample_Position", "Index1_ID", "Index1_Pos", "Index2_ID", "Index2_Pos")

  def Display_Error_Dialog(self, message: str):
    """The error shown to the operator ("Press OK to abort run"), then the run aborted."""
    self._ask("message", message, title="Press OK to abort run")
    raise RunAborted(message)

  def Read_InputWorklist_ExcelOrCSVFormat(self, path, template):
    """The samples of the input worklist: a CSV (`;` between fields, a header row naming
    WORKLIST_COLUMNS), or an Excel workbook's Sheet1 (read with openpyxl, `.xlsx`; Venus reads `.xls`
    through Jet, which openpyxl can't: give the workbook saved as `.xlsx` or as CSV). Rows without a
    sample position are skipped. Returns (samples, indexes used, index 1 positions, index 1 ids,
    index 2 positions, index 2 ids, the template's labware once per sample, the sample positions,
    "plate" or "tube" for the indexes: a plate when the first index 1 position is A1 to H12). The
    sample ids are kept (`self.sample_ids`) for the sample tracking.

    Kept as the library has it: "indexes used" is 2 whenever there are samples, as every row adds an
    index 2 entry, empty or not (single indexing is never reported); with no samples, the first
    index position is read from an empty list (Venus: a runtime error)."""
    import csv

    path = str(path)
    ext = path.rsplit(".", 1)[-1].lower() if "." in path else ""
    t = self.v.seq(template)
    t.set_current_position(1)
    lab = (t.current_item() or ("", ""))[0]
    # left out: the status-window line ("Read input worklist : <file>")
    try:
      if ext == "csv":
        with open(path, newline="", encoding="utf-8-sig") as f:
          rows = list(csv.DictReader(f, delimiter=";"))
      elif ext in ("xls", "xlsx"):
        rows = _sheet_rows(path, "Sheet1")
      else:
        rows = None
    except (OSError, ValueError, KeyError, ImportError) as e:  # ImportError: no openpyxl here
      kind = "CSV" if ext == "csv" else "XLS"
      self.Display_Error_Dialog(f"Protocol cannot open Input Worklist ({kind} format). Please review file structure "
                                f"and content. System will abort ({e})")
    if rows is None:
      self.Display_Error_Dialog("Protocol cannot open Input Worklist (Not supported format). Please review file "
                                "structure and content. System will abort")
    rows = [{k: ("" if r.get(k) is None else str(r.get(k))) for k in self.WORKLIST_COLUMNS} for r in rows]
    rows = [r for r in rows if r["Sample_Position"] != ""]
    self.sample_ids = [r["Sample_ID"] for r in rows]
    col = {k: [r[k] for r in rows] for k in self.WORKLIST_COLUMNS}
    if not rows:
      raise RunAborted("the input worklist holds no samples (Venus: the first index position of an empty list)")
    plate = {f"{r}{c}" for c in range(1, 13) for r in "ABCDEFGH"}
    index_format = "plate" if col["Index1_Pos"][0] in plate else "tube"
    used = 2 if col["Index2_Pos"] else (1 if col["Index1_Pos"] else 0)
    return (len(rows), used, col["Index1_Pos"], col["Index1_ID"], col["Index2_Pos"], col["Index2_ID"],
            [lab] * len(rows), col["Sample_Position"], index_format)

  def _ask(self, kind: str, text: str, **kw):
    if self.operator is None:
      raise RunAborted(f"the operator is asked, and there is none: {text}")
    return getattr(self.operator, kind)(text, **kw)

  def _carrier_sites(self, carrier: str) -> list[tuple[str, str]]:
    """The carrier's (site id, labware) holding labware, in the layout's order."""
    return [(p.fields.get("SiteId", ""), p.id) for p in self.v.layout.labware if p.fields.get("Template") == carrier]

  async def _unload(self, carrier: str) -> None:
    if self.star is not None:
      await self.star.autoload.unload_carrier(self.v.resource(carrier))

  async def _load(self, carrier: str, barcodes: str, report: bool = True) -> None:
    """LoadCarrier, reading the barcodes asked ("?" or "*": all; "TIP1,TIP2": those sites; "":
    none), and the report's rows for the barcodes read (_DATA_AnalyzeStepReturnValueOfAutoLoadStep...).
    Unverified (CONVENTIONS §11): that PLR's container positions are the carrier's sites holding
    labware in the layout's order, and the position Venus gives a plate's barcode (none here)."""
    from venus_interop.plr.libraries.ngsstar_report import AUTOLOAD_OK

    if self.star is None:
      return
    read = await self.star.autoload.load_carrier(self.v.resource(carrier), barcode_reading=barcodes != "")
    if not report:
      return
    wanted = None if barcodes in ("?", "*") else {p for p in barcodes.split(",") if p}
    sites = self._carrier_sites(carrier)
    for position, barcode in sorted((read or {}).get("container_barcodes", {}).items()):
      site, labware = sites[position] if position < len(sites) else ("", "")
      if barcode and (wanted is None or site in wanted):
        # the target and the pipetting columns are variables the library never assigns: "0" if one
        # is integer 0 (CONVENTIONS §11)
        self.report.add({**self.report.stamp(), "LabID Source": labware, "Barcode": barcode,
                         "Transport information": AUTOLOAD_OK, "LabID Target": "0", "PosID Target": "0",
                         "Aspiration information": "0", "Dispense information": "0"})

  async def Load_Carrier(self, ML_STAR, carrier: str, skip_loading=0, loadable=1, no_dialogs=0,
                         main_image: str = "", sub_image: str = "") -> bool:
    """Have the operator load `carrier` (one RunData names; else nothing is done): the autoload
    takes it out (both of a tip pair on an NGS STAR), the operator is shown what goes on it
    (positions, quantities, what) with its pictures, a tip carrier's counter is read and edited
    for its tip type, and the autoload loads it back, reading barcodes; the barcodes read go into
    the run report. `skip_loading` or not `loadable`: only the dialog and the tip counters.
    Kept as the library has it: of a tip pair's loads, only the last one's barcodes are reported."""
    if carrier not in self.carriers_to_load:  # ARRAYTOOLS::Lookup (inferred: PORTING)
      return False
    moved = not int(skip_loading) and int(loadable)
    pair = TIP_CARRIERS[50 if carrier == TIP50_01 else 300] if carrier in (TIP50_01, TIP300_01) else None
    await self._cover_unlocked()
    if moved:
      for c in (pair if pair and self.ngsstar else (carrier,)):
        await self._unload(c)
    table, plt1000, names, last = self._loading_list(carrier)
    if not int(no_dialogs):
      images = [self.pictures + main_image] + ([self.pictures + sub_image] if sub_image else [])
      if not self._ask("show_table", "Press OK to Continue, Cancel to Abort", title=carrier,
                       columns=["Position", "Quantity", "Description"], rows=table, images=images):
        raise RunAborted(f"{carrier}: loading cancelled")
      if "TIP_" in carrier or "Tip" in carrier:
        tip_type = 0  # never assigned when no row names a tip type (CONVENTIONS §11: integer 0)
        name = (last or {}).get("DB_ReagentName") or ""
        if TIP_NAMES[50] in name:
          tip_type = 1
        if TIP_NAMES[300] in name:
          tip_type = 2
        if any(TIP_NAMES[1000] in (n or "") for n in names):
          tip_type = 3
        units = (last or {}).get("DB_Units") or 0
        await self.DATA_TipCounterRead(ML_STAR, units, units, units, tip_type, 1, 0)
        self.DATA_TipCounterWrite(ML_STAR)
    if moved:
      await self._cover_locked()
      if pair:
        both = pair[:2 if self.ngsstar else 1]
        for c in both:
          await self._load(c, "?", report=c == both[-1])
      else:
        await self._load(carrier, plt1000 if carrier == PLT_TIP1000 else "*")
    return True

  # -- tip counters --------------------------------------------------------------------------

  def DATA_TipCounterWrite(self, ML_STAR=None) -> None:
    """The library's copies of the tip sequences into their counters (STAR_TIP_TOOLS::TipCounter_Write)."""
    for t, seq in self.tips.items():
      self.planner.TipCounter_Write(ML_STAR, seq, self.counters[t])

  async def DATA_TipCounterRead(self, ML_STAR, required50, required300, required1000, tip_types, always_edit=1,
                                show_calculation=0) -> None:
    """The tip sequences of `tip_types` (1: 50, 2: 300, 3: 1000 uL, 4: all) from their counters,
    the operator editing them first when `always_edit`. With fewer tips than required: past what
    the deck holds (960 of 50 or 300 uL, 192 of 1000), the operator is told a reload will come;
    otherwise shown what to refill, asked whether to refill every type, and each type to refill
    is unloaded, its counter edited, and loaded again (its barcodes into the report).

    Kept as the library has them: what is left is the end less the current position plus one, so
    a used-up rack (both 0) counts one tip; the operator's Cancel on the refill list is acted on
    only after the question that follows it, and ignored on the reload notice; the channel count
    the refill reads is never set in this library, so it takes the Clinical STARlet's branch and
    an NGS STAR refills only the first carrier of each pair. The consumed-tip counts the library
    works out are used nowhere and not ported."""
    which = {1: (50,), 2: (300,), 3: (1000,), 4: TIPS}.get(int(tip_types), ())
    required = dict(zip(TIPS, (int(required50), int(required300), int(required1000))))
    texts = dict.fromkeys(TIPS, NO_RELOAD)
    short = False
    for t in which:
      if int(always_edit):
        self.planner.TipCounter_Edit(ML_STAR, self.tips[t], self.counters[t])
        self.DATA_TipCounterWrite(ML_STAR)
      else:
        self.planner.TipCounter_Read(ML_STAR, self.tips[t], self.counters[t])
      s = self.tips[t]
      if s.end - s.cur + 1 < required[t]:
        short = True
        full = required[t] > (96 if t == 1000 else 960)
        texts[t] = ("\nPlease load 2 (only 1 for MPH version) full racks of 1000ul tips.  You will need to come back "
                    "later to replenish." if full and t == 1000 else
                    f"\nPlease load two full carriers of {t}ul tips.  You will need to come back later to replenish."
                    if full else f"\nPlease load more {t}ul tips.  You will need at least this many:  {required[t]}")
    if short:
      later = any(required[t] > (192 if t == 1000 else 960) for t in which)  # t_blnReduceTip1000sequence never set
      items = [f" Tip type {t} ul filter : {texts[t]}" for t in TIPS]
      scheduled = ("The amount of needed tips exceeds the maximum capacity of available tips. A reload step will be "
                   "scheduled within the process.")
      if later:
        self._ask("confirm", "Not enough tips available.", title=scheduled, items=items)
        texts = dict.fromkeys(TIPS, NO_RELOAD)
      else:
        ok = self._ask("confirm", "Not enough tips available - refill tips according to shown list", title=scheduled,
                       items=items)
        refill_all = self._ask("confirm", "Preventative tip exchange : Do you want to take the chance to reload all "
                                          "tip types?", yes_no=True)
        if not ok:
          raise RunAborted("not enough tips: cancelled")
        if refill_all:
          texts = dict.fromkeys(TIPS, PREVENTIVE)
      for t in TIPS:
        if texts[t] != NO_RELOAD:
          await self._refill(ML_STAR, t)
    self.DATA_TipCounterWrite(ML_STAR)

  async def _refill(self, ML_STAR, t: int) -> None:
    """A tip type's carrier out, its counter edited, and in again. The library unlocks and locks
    the cover around it (LockFrontCover only, no cover monitoring switched): here the operator is
    asked to close the cover while it is open, as the lock's retry does."""
    carriers = TIP_CARRIERS[t][:1]  # t_intNumberOfAllAvailableChannels == 8 never holds (see DATA_TipCounterRead)
    for c in carriers:
      await self._unload(c)
    self.planner.TipCounter_Edit(ML_STAR, self.tips[t], self.counters[t])
    self.DATA_TipCounterWrite(ML_STAR)
    for c in carriers:
      await self._load(c, "*")
    if self.lock_unlock and self.star is not None and self.star.front_cover is not None:
      while await self.star.front_cover.request_state() == "open":
        self._ask("message", COVER_OPEN, title="Front cover monitoring")


async def cover_unlocked(star) -> None:
  """The front cover's monitoring off (C0CD) and the cover unlocked. PLR has no cover lock (its
  FrontCover leaves lock and unlock out, unverified on hardware): only the monitoring is set."""
  if star is not None and getattr(star, "front_cover", None) is not None:
    await star.front_cover.disable_control()


async def cover_locked(star, ask) -> None:
  """The cover locked, the operator asked (`ask`, as `_ask`) to close it while it is open (Venus
  retries the lock after its dialog), then the monitoring on again (C0CE)."""
  if star is None or getattr(star, "front_cover", None) is None:
    return
  while await star.front_cover.request_state() == "open":
    ask("message", COVER_OPEN, title="Front cover monitoring")
  await star.front_cover.enable_control()


def _sheet_rows(path: str, sheet: str) -> list[dict]:
  """An Excel sheet's rows as dicts by its header row (openpyxl)."""
  from openpyxl import load_workbook

  ws = load_workbook(path, read_only=True, data_only=True)[sheet]
  values = list(ws.iter_rows(values_only=True))
  if not values:
    return []
  header = [str(h) if h is not None else "" for h in values[0]]
  return [dict(zip(header, r)) for r in values[1:]]


def _position_text(p: str) -> str:
  """_SortElements' position: 1 to 9 as 01 to 09; a list ("1,12") one a line, single digits so
  ("01" and "12" on two lines)."""
  if p in [str(i) for i in range(1, 10)]:
    return "0" + p
  return "\n".join("0" + t if len(t) == 1 else t for t in p.split(",") if t)
